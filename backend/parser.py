import os
from pathlib import Path
from typing import List, Dict, Any
import io

class DocumentParser:
    """
    Universal document parser that ingests and extracts structured text
    from multiple document types: PDF, DOCX, PPTX, XLSX, CSV, TXT, MD, JSON, etc.
    """

    @classmethod
    def parse_file(cls, file_path: Path, filename: str) -> List[Dict[str, Any]]:
        """
        Parses a file and returns a list of section dicts:
        [{ "text": "...", "page": 1, "section": "..." }]
        """
        ext = file_path.suffix.lower()
        
        if ext == ".pdf":
            return cls._parse_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls._parse_docx(file_path)
        elif ext == ".pptx":
            return cls._parse_pptx(file_path)
        elif ext in [".xlsx", ".xls"]:
            return cls._parse_excel(file_path)
        elif ext in [".csv", ".tsv"]:
            return cls._parse_csv(file_path, delimiter="," if ext == ".csv" else "\t")
        elif ext in [".json", ".jsonl"]:
            return cls._parse_json(file_path)
        else:
            # Fallback for text-based formats: .txt, .md, .py, .js, .html, .xml, .yaml, .yml, .log, etc.
            return cls._parse_text(file_path)

    @classmethod
    def _parse_pdf(cls, file_path: Path) -> List[Dict[str, Any]]:
        sections = []
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text()
                if text and text.strip():
                    sections.append({
                        "text": text.strip(),
                        "page": page_num + 1,
                        "section": f"Page {page_num + 1}"
                    })
            doc.close()
        except Exception as e:
            # In case PyMuPDF fails or encrypted
            sections.append({
                "text": f"Error parsing PDF: {str(e)}",
                "page": 1,
                "section": "Error"
            })
        return sections

    @classmethod
    def _parse_docx(cls, file_path: Path) -> List[Dict[str, Any]]:
        sections = []
        try:
            import docx
            doc = docx.Document(file_path)
            full_text = []
            for i, p in enumerate(doc.paragraphs):
                if p.text.strip():
                    full_text.append(p.text.strip())
            
            # Also extract tables
            for t_idx, table in enumerate(doc.tables):
                table_lines = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        table_lines.append(" | ".join(row_data))
                if table_lines:
                    full_text.append(f"\n[Table {t_idx + 1}]\n" + "\n".join(table_lines))

            if full_text:
                sections.append({
                    "text": "\n\n".join(full_text),
                    "page": 1,
                    "section": "Full Document"
                })
        except Exception as e:
            sections.append({
                "text": f"Error parsing DOCX: {str(e)}",
                "page": 1,
                "section": "Error"
            })
        return sections

    @classmethod
    def _parse_pptx(cls, file_path: Path) -> List[Dict[str, Any]]:
        sections = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for slide_num, slide in enumerate(prs.slides, start=1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    sections.append({
                        "text": "\n".join(slide_texts),
                        "page": slide_num,
                        "section": f"Slide {slide_num}"
                    })
        except Exception as e:
            sections.append({
                "text": f"Error parsing PPTX: {str(e)}",
                "page": 1,
                "section": "Error"
            })
        return sections

    @classmethod
    def _parse_excel(cls, file_path: Path) -> List[Dict[str, Any]]:
        sections = []
        try:
            import pandas as pd
            excel_file = pd.ExcelFile(file_path)
            for sheet_name in excel_file.sheet_names:
                df = pd.read_excel(file_path, sheet_name=sheet_name)
                # Convert first 500 rows to markdown/string table summary
                sheet_text = f"Sheet: {sheet_name}\nColumns: {', '.join(df.columns.astype(str))}\n\n"
                sheet_text += df.head(500).to_string(index=False)
                sections.append({
                    "text": sheet_text,
                    "page": 1,
                    "section": f"Sheet: {sheet_name}"
                })
        except Exception as e:
            sections.append({
                "text": f"Error parsing Excel: {str(e)}",
                "page": 1,
                "section": "Error"
            })
        return sections

    @classmethod
    def _parse_csv(cls, file_path: Path, delimiter: str = ",") -> List[Dict[str, Any]]:
        sections = []
        try:
            import pandas as pd
            df = pd.read_csv(file_path, sep=delimiter)
            csv_text = f"Columns: {', '.join(df.columns.astype(str))}\nTotal Rows: {len(df)}\n\n"
            csv_text += df.head(500).to_string(index=False)
            sections.append({
                "text": csv_text,
                "page": 1,
                "section": "Data Table"
            })
        except Exception as e:
            # Fallback to direct reading
            return cls._parse_text(file_path)
        return sections

    @classmethod
    def _parse_json(cls, file_path: Path) -> List[Dict[str, Any]]:
        try:
            import json
            raw_text = cls._read_file_content(file_path)
            data = json.loads(raw_text)
            formatted = json.dumps(data, indent=2)
            return [{
                "text": formatted,
                "page": 1,
                "section": "JSON Structure"
            }]
        except Exception:
            return cls._parse_text(file_path)

    @classmethod
    def _parse_text(cls, file_path: Path) -> List[Dict[str, Any]]:
        content = cls._read_file_content(file_path)
        return [{
            "text": content,
            "page": 1,
            "section": "Document Content"
        }]

    @staticmethod
    def _read_file_content(file_path: Path) -> str:
        for enc in ["utf-8", "utf-8-sig", "latin-1", "cp1252", "iso-8859-1"]:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except UnicodeDecodeError:
                continue
        # Fallback with ignore errors
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
